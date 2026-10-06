import re
import shutil
import sqlite3
import sys
import uuid
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile

from lxml import etree
from openpyxl import load_workbook


if len(sys.argv) != 3:
    raise SystemExit("Usage: import-team-history.py <database.sqlite> <source.xlsx>")

database_path = Path(sys.argv[1])
workbook_path = Path(sys.argv[2])
backup_path = database_path.with_name(
    f"{database_path.stem}.pre-team-history-{datetime.now():%Y%m%d-%H%M%S}.sqlite"
)
shutil.copy2(database_path, backup_path)

SHEET_MEMBERS = {
    "Muthukrishanan": "MHS299",
    "Rahul R": "MHS316",
    "Rahul B": "MHS382",
    "Kamesh": "MHS090",
    "Saslivikram": "MHS386",
    "Varadharaj": "MHS171",
    "Rahul": "MHS250",
    "Basha": "MHS392",
    "Micheal": "MHS281",
    "Karthik": "MHS278",
}


def clean(value):
    return str(value or "").strip()


def iso_date(value):
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = clean(value)
    for fmt in ("%d-%b-%y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            pass
    return ""


def number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def task_status(value):
    text = clean(value).lower()
    if text == "done":
        return "done"
    if "progress" in text:
        return "progress"
    if "hold" in text:
        return "hold"
    return "todo"


def priority(value):
    text = clean(value).lower()
    if "high" in text:
        return "high"
    if "low" in text:
        return "low"
    return "normal"


def classify(sheet, title, notes):
    text = f"{title} {notes}".lower()
    if sheet == "Muthukrishanan":
        if any(k in text for k in ("longform", "long form", "youtube", "simulive")):
            return "Longform"
        if any(k in text for k in (" ad ", "ads", "testimonial")):
            return "Ads"
        if "reel" in text:
            return "Reels"
    elif sheet == "Rahul R":
        if "carousel" in text:
            return "Carousel"
        if any(k in text for k in ("publish", "schedule", "meta", "comment", "data feed")):
            return "Schedule"
        if " ad" in text or "ads" in text:
            return "Ads"
    elif sheet == "Rahul B":
        if any(k in text for k in ("video carousel", "bytes")):
            return "Video Carousel"
        if " ai " in f" {text} " or "chatgpt" in text:
            return "AI Hours"
    elif sheet == "Kamesh":
        if any(k in text for k in ("shoot", "filming")):
            return "Shoot"
        if any(k in text for k in ("longform", "long form", "trailer", "youtube", " yt ")):
            return "Longform"
        if " ad" in text or "ads" in text:
            return "Ads"
        if "reel" in text:
            return "Reels"
    elif sheet == "Saslivikram":
        if any(k in text for k in ("longform", "long form", "youtube")):
            return "Longform"
        if " ad" in text or "ads" in text:
            return "Ads"
        if "reel" in text:
            return "Reels"
    elif sheet == "Varadharaj":
        if any(k in text for k in ("longform", "long form", "youtube", "simulive")):
            return "Longform"
        if " ai " in f" {text} " or "chatgpt" in text:
            return "AI Hours"
        if " ad" in text or "ads" in text:
            return "Ads"
    elif sheet == "Basha":
        if "story" in text:
            return "Story"
        if any(k in text for k in ("publish", "schedule", "comment deletion")):
            return "Schedule"
        if "carousel" in text:
            return "Carousel"
    elif sheet == "Micheal":
        if any(k in text for k in ("shoot", "filming")):
            return "Shoot"
        if " ai " in f" {text} " or "chatgpt" in text:
            return "AI Hours"
        if any(k in text for k in (" ad", "wellness", "insulin", "sweet", "phone call", "google", "reverse", "walking")):
            return "Ads"
    elif sheet == "Karthik":
        if "carousel" in text:
            return "Carousel"
        if "reel" in text:
            return "Reels"
    return "Others"


def hyperlink_map(path):
    """Return every exported hyperlink grouped by sheet and cell reference."""
    result = defaultdict(lambda: defaultdict(list))
    with ZipFile(path) as archive:
        main_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        workbook = etree.fromstring(archive.read("xl/workbook.xml"))
        relationships = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rel_targets = {node.get("Id"): node.get("Target") for node in relationships}
        ns = {"m": main_ns, "r": rel_ns}
        for sheet in workbook.xpath("//m:sheets/m:sheet", namespaces=ns):
            title = sheet.get("name")
            target = rel_targets[sheet.get(f"{{{rel_ns}}}id")]
            target = target.lstrip("/")
            if not target.startswith("xl/"):
                target = f"xl/{target}"
            target = target.replace("xl/xl/", "xl/")
            sheet_number = target.rsplit("sheet", 1)[-1].split(".", 1)[0]
            relation_path = f"xl/worksheets/_rels/sheet{sheet_number}.xml.rels"
            if relation_path not in archive.namelist():
                continue
            relation_doc = etree.fromstring(archive.read(relation_path))
            urls = {node.get("Id"): node.get("Target") for node in relation_doc}
            sheet_doc = etree.fromstring(archive.read(target))
            for node in sheet_doc.xpath("//m:hyperlink", namespaces=ns):
                cell = node.get("ref")
                rel_id = node.get(f"{{{rel_ns}}}id")
                url = urls.get(rel_id, "")
                if cell and url and url not in result[title][cell]:
                    result[title][cell].append(url)
    return result


workbook = load_workbook(workbook_path, data_only=True)
links_by_cell = hyperlink_map(workbook_path)
connection = sqlite3.connect(database_path)
connection.row_factory = sqlite3.Row

people = {
    clean(row["employee_id"]).upper(): row
    for row in connection.execute("SELECT * FROM people WHERE employee_id IS NOT NULL")
}
plans = defaultdict(dict)
for row in connection.execute(
    """
    SELECT p.*, g.name AS goal_name
    FROM plans p JOIN goal_types g ON g.id=p.type_id
    WHERE p.month='2026-09'
    """
):
    plans[row["member_id"]][row["goal_name"]] = row

existing_entries = Counter(
    (
        row["member_id"],
        row["work_date"],
        clean(row["title"]).lower(),
        clean(row["url"]),
    )
    for row in connection.execute(
        """
        SELECT d.member_id,d.work_date,d.title,s.url
        FROM daily_tasks d LEFT JOIN submissions s ON s.daily_task_id=d.id
        """
    )
)

inserted_tasks = Counter()
inserted_submissions = Counter()
skipped = Counter()
categories = Counter()

try:
    with connection:
        for sheet_name, employee_id in SHEET_MEMBERS.items():
            # Rahul's historical work was imported separately because rows with
            # multiple outputs use output-specific titles and links.
            if sheet_name == "Rahul":
                skipped["Rahul (managed separately)"] += 1
                continue
            person = people.get(employee_id)
            if not person:
                raise RuntimeError(f"No employee profile for {employee_id} ({sheet_name})")
            sheet = workbook[sheet_name]
            member_plans = plans[person["id"]]
            if "Others" not in member_plans:
                raise RuntimeError(f"No Others plan for {person['name']}")

            for row_number in range(6, sheet.max_row + 1):
                work_date = iso_date(sheet.cell(row_number, 1).value)
                title = clean(sheet.cell(row_number, 2).value)
                if not work_date or work_date < "2026-09-15" or not title:
                    continue
                source_status = task_status(sheet.cell(row_number, 6).value)
                due_date = iso_date(sheet.cell(row_number, 4).value) or work_date
                notes_value = clean(sheet.cell(row_number, 9).value)
                category = classify(sheet_name, title, notes_value)
                plan = member_plans.get(category) or member_plans["Others"]
                category = plan["goal_name"]
                urls = links_by_cell[sheet_name].get(f"I{row_number}", [])
                if not urls:
                    cell_link = sheet.cell(row_number, 9).hyperlink
                    if cell_link and cell_link.target:
                        urls = [cell_link.target]
                outputs = urls if source_status == "done" and urls else ([""] if source_status == "done" else [None])

                for output_index, url in enumerate(outputs, start=1):
                    task_title = title if output_index == 1 else f"{title} — Output {output_index}"
                    entry_key = (person["id"], work_date, task_title.lower(), clean(url))
                    if existing_entries[entry_key] > 0:
                        existing_entries[entry_key] -= 1
                        skipped[person["name"]] += 1
                        continue

                    task_id = str(uuid.uuid4())
                    estimated = number(plan["hours_per_job"])
                    actual = number(sheet.cell(row_number, 8).value) if output_index == 1 else 0
                    note_parts = ["Historical work entry"]
                    plain_note = re.sub(r"https?://\S+", "", notes_value).strip()
                    if plain_note and plain_note != title:
                        note_parts.append(plain_note)
                    submitted_at = f"{work_date}T18:00:00.000+05:30" if source_status == "done" else None
                    connection.execute(
                        """
                        INSERT INTO daily_tasks(
                          id,member_id,work_date,title,category,due_date,priority,status,
                          estimated_hours,actual_hours,notes,version,started_at,
                          elapsed_seconds,submitted_at,original_work_date
                        ) VALUES(?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)
                        """,
                        (
                            task_id, person["id"], work_date, task_title, category, due_date,
                            priority(sheet.cell(row_number, 5).value), source_status, estimated,
                            actual, "\n".join(note_parts), 1, None, 0, submitted_at, work_date,
                        ),
                    )
                    inserted_tasks[person["name"]] += 1
                    categories[(person["name"], category)] += 1

                    if source_status == "done":
                        connection.execute(
                            """
                            INSERT INTO submissions(
                              id,plan_id,title,url,completed,quantity,status,feedback,version,daily_task_id
                            ) VALUES(?,?,?,?,?,1,'review','',1,?)
                            """,
                            (str(uuid.uuid4()), plan["id"], task_title, url or "", work_date, task_id),
                        )
                        inserted_submissions[person["name"]] += 1
finally:
    connection.close()

print(f"backup={backup_path}")
print("tasks=" + ", ".join(f"{name}:{count}" for name, count in inserted_tasks.items()))
print("submissions=" + ", ".join(f"{name}:{count}" for name, count in inserted_submissions.items()))
print("skipped=" + ", ".join(f"{name}:{count}" for name, count in skipped.items()))
for (name, category), count in sorted(categories.items()):
    print(f"category={name}|{category}|{count}")
